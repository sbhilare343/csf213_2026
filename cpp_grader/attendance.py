from __future__ import annotations
from pathlib import Path
from typing import Any
import openpyxl
from .config import AppConfig
from .errors import InputValidationError
from .models import Student
from .utils import cell_text, norm

ALIASES = {
    "student_id": ("ID number", "Student ID", "Campus ID", "ID"),
    "name": ("Name", "Student Name"),
    "email": ("Email address", "Email"),
    "groups": ("Groups", "Group"),
    "zone": ("Zone",),
    "seat": ("Seat",),
    "absent": ("Absent",),
    "makeup": ("Makeup", "Make-up"),
    "valid_submission": ("Valid submission on Quanta", "Valid submission", "Submission"),
}

def _find_columns(headers: tuple[Any, ...], configured: dict[str, Any]) -> dict[str, int]:
    hmap = {norm(h): i for i, h in enumerate(headers) if h is not None and str(h).strip()}
    out: dict[str, int] = {}
    for field, aliases in ALIASES.items():
        preferred = configured.get(field)
        candidates = ([str(preferred)] if preferred else []) + list(aliases)
        for candidate in candidates:
            if norm(candidate) in hmap:
                out[field] = hmap[norm(candidate)]
                break
    missing = [field for field in ("student_id", "name") if field not in out]
    if missing:
        raise InputValidationError(f"Attendance sheet is missing required columns {missing}; found {list(headers)}")
    return out

def as_bool(value: Any) -> bool | None:
    if isinstance(value, bool): return value
    if value is None or str(value).strip() == "": return None
    s = str(value).strip().casefold()
    if s in {"true", "yes", "y", "1", "absent", "granted"}: return True
    if s in {"false", "no", "n", "0", "present", "na", "n/a"}: return False
    return None

def load_roster(config: AppConfig) -> list[Student]:
    path = config.path_value("paths", "attendance_xlsx")
    if not path or not path.is_file():
        raise InputValidationError(f"Attendance workbook is missing: {path}")
    cfg = config.section("attendance")
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise InputValidationError(f"Cannot open attendance workbook {path}: {exc}") from exc
    try:
        sheet_name = cfg.get("sheet_name")
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise InputValidationError(f"Attendance sheet {sheet_name!r} not found; sheets={wb.sheetnames}")
            ws = wb[sheet_name]
        else:
            ws = wb.active
        header_row = int(cfg.get("header_row", 1))
        header = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), None)
        if header is None:
            raise InputValidationError("Attendance worksheet is empty.")
        positions = _find_columns(header, cfg.get("columns", {}))
        students: list[Student] = []
        seen: dict[str, int] = {}
        for row_index, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            def get(field: str) -> Any:
                idx = positions.get(field)
                return row[idx] if idx is not None and idx < len(row) else None
            student_id = cell_text(get("student_id"))
            name = cell_text(get("name"))
            if not student_id and not name:
                continue
            if not student_id or not name:
                raise InputValidationError(f"Attendance row {row_index} must have both a student ID and name.")
            key = norm(student_id)
            if key in seen:
                raise InputValidationError(f"Duplicate student ID {student_id!r} in attendance rows {seen[key]} and {row_index}.")
            seen[key] = row_index
            absent = as_bool(get("absent"))
            makeup = cell_text(get("makeup"))
            if makeup.casefold() == "granted":
                attendance = "Makeup"
            elif absent is True:
                attendance = "Absent"
            else:
                attendance = "Present"
            students.append(Student(
                student_id=student_id, name=name, email=cell_text(get("email")),
                groups=cell_text(get("groups")), zone=cell_text(get("zone")), seat=cell_text(get("seat")),
                attendance=attendance, absent=absent, makeup=makeup,
                valid_submission=as_bool(get("valid_submission")), roster_row=row_index,
            ))
    finally:
        wb.close()
    if not students:
        raise InputValidationError("No student rows were found in the attendance workbook.")
    return students
