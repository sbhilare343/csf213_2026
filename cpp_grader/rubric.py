from __future__ import annotations
from pathlib import Path
from typing import Any
import re
import math
import openpyxl
import fitz
from .case_catalog import SUPPORTED_TEST_IDS
from .config import AppConfig
from .errors import InputValidationError
from .models import TestCase
from .utils import cell_text, norm, json_write

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "id": ("ID", "Test ID", "TestCase ID", "Test Case ID"),
    "method": ("UUT (method)", "Method", "UUT"),
    "section": ("Section",),
    "category": ("Category",),
    "scenario": ("Scenario", "Description"),
    "drivers_added": ("Drivers added to RideService (position: rating)", "Drivers Added", "Setup"),
    "steps": ("Steps (in order)", "Steps"),
    "expected": ("Expected result", "Expected"),
    "catches": ("Bug this test catches", "Bug Caught", "Purpose"),
    "marks": ("Marks", "Points", "Maximum Marks"),
}


def _header_positions(headers: list[Any]) -> dict[str, int]:
    mapped = {norm(h): i for i, h in enumerate(headers) if h is not None and str(h).strip()}
    positions: dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if norm(alias) in mapped:
                positions[field] = mapped[norm(alias)]
                break
    missing = [x for x in ("id", "scenario", "steps", "expected", "marks") if x not in positions]
    if missing:
        raise InputValidationError(
            f"Rubric worksheet is missing columns for {', '.join(missing)}. Headers found: "
            + repr([str(x) for x in headers if x is not None])
        )
    return positions


def load_rubric(config: AppConfig) -> list[TestCase]:
    path = config.path_value("paths", "rubric_xlsx")
    if not path or not path.is_file():
        raise InputValidationError(f"Official rubric workbook is missing: {path}")
    rubric_cfg = config.section("rubric")
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise InputValidationError(f"Could not open rubric workbook {path}: {exc}") from exc
    try:
        sheet_name = rubric_cfg.get("sheet_name")
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise InputValidationError(f"Rubric sheet {sheet_name!r} not found. Available: {workbook.sheetnames}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        header_row = int(rubric_cfg.get("header_row", 1))
        iterator = sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
        header = next(iterator, None)
        if header is None:
            raise InputValidationError(f"Rubric worksheet {sheet.title!r} is empty.")
        positions = _header_positions(list(header))
        cases: list[TestCase] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            raw_id = row[positions["id"]] if positions["id"] < len(row) else None
            case_id = cell_text(raw_id)
            if not case_id:
                continue
            def val(field: str) -> str:
                idx = positions.get(field)
                return cell_text(row[idx]) if idx is not None and idx < len(row) else ""
            try:
                marks = float(row[positions["marks"]])
            except (ValueError, TypeError, IndexError) as exc:
                raise InputValidationError(f"Test case {case_id} at row {row_number} has invalid marks.") from exc
            if not math.isfinite(marks) or marks < 0:
                raise InputValidationError(f"Test case {case_id} at row {row_number} must have finite, non-negative marks.")
            if not marks.is_integer():
                marks = round(marks, 4)
            cases.append(TestCase(
                id=case_id, method=val("method"), section=val("section"), category=val("category"),
                scenario=val("scenario"), drivers_added=val("drivers_added"), steps=val("steps"),
                expected=val("expected"), catches=val("catches"), marks=marks, source_row=row_number,
            ))
    finally:
        workbook.close()

    ids = [c.id for c in cases]
    if len(ids) != len(set(x.casefold() for x in ids)):
        duplicates = sorted({x for x in ids if sum(y.casefold() == x.casefold() for y in ids) > 1})
        raise InputValidationError(f"Duplicate testcase IDs in workbook: {duplicates}")
    unsupported = sorted(set(ids) - set(SUPPORTED_TEST_IDS))
    missing = sorted(set(SUPPORTED_TEST_IDS) - set(ids))
    if unsupported or missing:
        raise InputValidationError(
            "Workbook test IDs do not match the implemented C++ harness. "
            f"Missing IDs: {missing or 'none'}; unsupported IDs: {unsupported or 'none'}. "
            "Never grade with only a partial/mismatched rubric."
        )
    expected_count = int(rubric_cfg.get("expected_testcase_count", len(SUPPORTED_TEST_IDS)))
    if len(cases) != expected_count:
        raise InputValidationError(f"Expected {expected_count} test cases but found {len(cases)}.")
    total = round(sum(c.marks for c in cases), 4)
    expected_total = float(rubric_cfg.get("expected_total_marks", total))
    if not math.isfinite(expected_total) or expected_total < 0:
        raise InputValidationError("rubric.expected_total_marks must be finite and non-negative.")
    if abs(total - expected_total) > 1e-8:
        raise InputValidationError(f"Rubric total is {total}, expected {expected_total}; resolve this before grading.")
    return sorted(cases, key=lambda c: SUPPORTED_TEST_IDS.index(c.id))


def inspect_pdf(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {"exists": False, "path": str(path), "error": "not found"}
    try:
        with fitz.open(path) as doc:
            parts = [page.get_text("text") for page in doc]
            full = "\n".join(parts)
            ids_present = [case_id for case_id in SUPPORTED_TEST_IDS if re.search(rf"(?<![A-Za-z0-9]){re.escape(case_id)}(?![A-Za-z0-9])", full)]
            return {
                "exists": True, "path": str(path), "page_count": len(doc),
                "extractable_characters": len(full), "testcase_ids_found": ids_present,
                "testcase_ids_missing": sorted(set(SUPPORTED_TEST_IDS) - set(ids_present)),
                "mentions_60_marks": bool(re.search(r"\b60\b", full)),
                "text_excerpt": full[:1200],
            }
    except Exception as exc:
        return {"exists": True, "path": str(path), "error": f"{type(exc).__name__}: {exc}"}


def validate_pdf_reconciliation(path: Path) -> list[str]:
    info = inspect_pdf(path)
    if not info.get("exists"):
        raise InputValidationError(f"Official testcase PDF is missing: {path}")
    if info.get("error"):
        raise InputValidationError(f"Could not read official testcase PDF: {info['error']}")
    warnings = []
    if info.get("testcase_ids_missing"):
        warnings.append("PDF text extraction did not locate these test IDs (inspect PDF visually): " + ", ".join(info["testcase_ids_missing"]))
    if info.get("extractable_characters", 0) < 100:
        warnings.append("PDF contains little extractable text; visual review is required.")
    return warnings


def export_rubric_catalog(path: Path, cases: list[TestCase], pdf_info: dict[str, Any] | None = None) -> None:
    payload = {
        "assignment": "CS F213 OOP Lab Midsem - C++",
        "source_workbook": "Grader/testcases.xlsx",
        "source_pdf": "Grader/grader-testcases.pdf",
        "maximum_marks": round(sum(c.marks for c in cases), 4),
        "testcase_count": len(cases),
        "pdf_audit": pdf_info or {},
        "testcases": [
            {"id": c.id, "method": c.method, "section": c.section, "category": c.category,
             "scenario": c.scenario, "drivers_added": c.drivers_added, "steps": c.steps,
             "expected": c.expected, "catches": c.catches, "marks": c.marks,
             "source_row": c.source_row}
            for c in cases
        ],
    }
    json_write(path, payload)
