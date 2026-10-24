from __future__ import annotations
import csv
import json
from pathlib import Path
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from .models import Student, Submission, TestCase
from .utils import json_write, norm

NAVY = "17365D"
BLUE = "D9EAF7"
LIGHT_RED = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRAY = "666666"


def _style_sheet(ws, widths: dict[int, float] | None = None) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 else None
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    thin = Side(style="hair", color="D9E2F3")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color="222222")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    if widths:
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
    else:
        for col in range(1, ws.max_column + 1):
            max_len = max((len(str(ws.cell(row, col).value or "")) for row in range(1, min(ws.max_row, 150) + 1)), default=10)
            ws.column_dimensions[get_column_letter(col)].width = min(36, max(10, max_len + 2))


def _student_key(r: dict[str, Any]) -> str:
    return norm(r.get("student_id"))


def _gradebook_row(r: dict[str, Any], cases: list[TestCase]) -> dict[str, Any]:
    case_map = {c.get("testcase_id"): c for c in r.get("case_results", [])}
    row: dict[str, Any] = {
        "Student ID": r.get("student_id", ""), "Student Name": r.get("name", ""),
        "Attendance Status": r.get("attendance", ""), "Absent Flag (sheet)": r.get("absent"),
        "Makeup (sheet)": r.get("makeup", ""),
        "Zone": r.get("zone", ""), "Seat": r.get("seat", ""),
        "Valid submission on Quanta": r.get("valid_submission"),
        "Submission Status": r.get("submission_status", ""),
        "Matched By": r.get("matched_by", ""), "Source Files": r.get("source_count", 0),
        "Compilation Status": r.get("compilation_status", ""),
    }
    for c in cases:
        item = case_map.get(c.id)
        row[f"{c.id} ({c.marks:g})"] = item.get("earned_marks") if item else None
    row["Total"] = r.get("earned_marks")
    row["Out Of"] = r.get("max_marks")
    row["Percentage"] = r.get("percentage")
    row["Review Flag"] = r.get("review_flag", "")
    return row


def write_grade_reports(output_dir: Path, results: list[dict[str, Any]], cases: list[TestCase],
                        unmatched: list[Submission], summary: dict[str, Any], *,
                        xlsx_name: str = "grades.xlsx", csv_name: str = "grades.csv") -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    sorted_results = sorted(results, key=lambda r: (str(r.get("name", "")).casefold(), str(r.get("student_id", ""))))
    rows = [_gradebook_row(r, cases) for r in sorted_results]
    columns = list(rows[0].keys()) if rows else ["Student ID", "Student Name", "Attendance Status", "Absent Flag (sheet)", "Makeup (sheet)", "Submission Status"]

    csv_path = output_dir / csv_name
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: ("" if v is None else v) for k, v in row.items()})

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Gradebook"
    ws.append(columns)
    for row in rows: ws.append([row.get(c) for c in columns])
    _style_sheet(ws, {1: 20, 2: 30, 3: 16, 4: 18, 5: 16, 6: 10, 7: 10, 8: 24, 9: 24, 10: 18, 11: 12, 12: 18})
    first_case_col = next((i for i, name in enumerate(columns, start=1) if name.startswith(f"{cases[0].id} (")), len(columns) + 1) if cases else len(columns) + 1
    for col in range(first_case_col, ws.max_column + 1): ws.column_dimensions[get_column_letter(col)].width = 10
    for row_num in range(2, ws.max_row + 1):
        for col_num in range(first_case_col, first_case_col + len(cases)):
            ws.cell(row_num, col_num).number_format = "0.##"
        total_col = columns.index("Total") + 1
        out_col = columns.index("Out Of") + 1
        pct_col = columns.index("Percentage") + 1
        ws.cell(row_num, total_col).number_format = "0.##"
        ws.cell(row_num, out_col).number_format = "0.##"
        ws.cell(row_num, pct_col).number_format = "0.00"  # percentage values are stored on a 0-100 scale
        status_col = columns.index("Compilation Status") + 1
        if str(ws.cell(row_num, status_col).value).upper() == "FAILED":
            ws.cell(row_num, status_col).fill = PatternFill("solid", fgColor=LIGHT_RED)
        if ws.cell(row_num, total_col).value is None:
            ws.cell(row_num, total_col).fill = PatternFill("solid", fgColor=LIGHT_YELLOW)

    details = wb.create_sheet("Testcase Details")
    detail_headers = ["Student ID", "Student Name", "Test ID", "Test Description", "Status", "Earned Marks", "Maximum Marks", "Duration (s)", "Message", "Stdout (truncated)", "Stderr (truncated)"]
    details.append(detail_headers)
    for r in sorted_results:
        for c in r.get("case_results", []):
            details.append([r.get("student_id"), r.get("name"), c.get("testcase_id"), c.get("testcase_name"), c.get("status"),
                            c.get("earned_marks"), c.get("max_marks"), c.get("duration_seconds"), c.get("message"), c.get("stdout"), c.get("stderr")])
    _style_sheet(details, {1: 20, 2: 30, 3: 12, 4: 46, 5: 23, 6: 13, 7: 14, 8: 14, 9: 52, 10: 50, 11: 50})

    exceptions = wb.create_sheet("Exceptions")
    exceptions.append(["Student ID", "Student Name", "Attendance Status", "Absent Flag (sheet)", "Submission Status", "Compilation Status", "Earned Marks", "Review Flag"])
    for r in sorted_results:
        if r.get("review_flag") or r.get("earned_marks") is None:
            exceptions.append([r.get("student_id"), r.get("name"), r.get("attendance"), r.get("absent"),
                              r.get("submission_status"), r.get("compilation_status"), r.get("earned_marks"), r.get("review_flag")])
    for sub in unmatched:
        exceptions.append(["", sub.folder_name, "", None, sub.status, "NOT_RUN", None, sub.issue])
    _style_sheet(exceptions, {1: 20, 2: 30, 3: 16, 4: 18, 5: 24, 6: 20, 7: 14, 8: 80})

    summ = wb.create_sheet("Run Summary")
    summ.append(["Field", "Value"])
    for key, value in summary.items():
        if isinstance(value, (str, int, float, bool)) or value is None:
            summ.append([key, value])
        else:
            summ.append([key, json.dumps(value, ensure_ascii=False)])
    _style_sheet(summ, {1: 38, 2: 100})
    xlsx_path = output_dir / xlsx_name
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    return {"xlsx": xlsx_path, "csv": csv_path}


def write_reconciliation(path: Path, roster: list[Student], matches: dict[str, list[Submission]], unmatched: list[Submission]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    roster_ws = wb.active; roster_ws.title = "Roster Reconciliation"
    roster_ws.append(["Student ID", "Name", "Attendance Status", "Absent Flag (sheet)", "Makeup Flag (sheet)", "Zone", "Seat", "Valid submission on Quanta", "Submission Match", "Submission Folder", "C++ Source Files", "Matched By", "Issue"])
    for student in sorted(roster, key=lambda s: s.name.casefold()):
        subs = matches.get(norm(student.student_id), [])
        if not subs:
            roster_ws.append([student.student_id, student.name, student.attendance, student.absent, student.makeup, student.zone, student.seat,
                              student.valid_submission, "NO_SUBMISSION", "", 0, "", "No candidate directory matched this student."])
        elif len(subs) > 1:
            roster_ws.append([student.student_id, student.name, student.attendance, student.absent, student.makeup, student.zone, student.seat,
                              student.valid_submission, "DUPLICATE_SUBMISSIONS", "; ".join(s.folder_name for s in subs),
                              sum(len(s.source_files) for s in subs), "", "Multiple candidate directories; manual review required."])
        else:
            sub = subs[0]
            roster_ws.append([student.student_id, student.name, student.attendance, student.absent, student.makeup, student.zone, student.seat,
                              student.valid_submission, sub.status, sub.folder_name, len(sub.source_files), sub.matched_by, sub.issue])
    _style_sheet(roster_ws, {1: 20, 2: 30, 3: 16, 4: 18, 5: 18, 6: 10, 7: 10, 8: 25, 9: 24, 10: 42, 11: 16, 12: 22, 13: 70})

    unmatched_ws = wb.create_sheet("Unmatched Folders")
    unmatched_ws.append(["Folder", "Status", "C++ Sources", "Headers", "Issue"])
    for sub in unmatched:
        unmatched_ws.append([sub.folder_name, sub.status, len(sub.source_files), len(sub.header_files), sub.issue])
    _style_sheet(unmatched_ws, {1: 45, 2: 24, 3: 16, 4: 12, 5: 80})
    wb.save(path)
